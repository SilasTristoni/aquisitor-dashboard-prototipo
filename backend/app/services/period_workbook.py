from __future__ import annotations

import csv
import io
import json
from bisect import bisect_left
from copy import copy
from datetime import datetime
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.version import APPLICATION_VERSION
from app.schemas.contracts import PeriodReportRequest
from app.services.period_documents import (
    CHANNEL_COLORS,
    POWER_COLOR,
    format_duration_pt,
    padded_bounds,
)

NAVY = "17233F"
BLUE = "2563EB"
PALE_BLUE = "EAF0FF"
PALE_GRAY = "F3F6FB"
WHITE = "FFFFFF"


def _local_excel(value: datetime, zone: ZoneInfo) -> datetime:
    localized = value.astimezone(zone)
    return localized.replace(tzinfo=None)


def _header(sheet: Any, row: int, columns: int) -> None:
    for cell in sheet[row][:columns]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center")


def _finish_sheet(sheet: Any, freeze: str = "A2") -> None:
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, sheet.max_column + 1):
        values = [str(sheet.cell(row, column).value or "") for row in range(1, sheet.max_row + 1)]
        sheet.column_dimensions[get_column_letter(column)].width = min(
            34, max(11, max(map(len, values), default=0) + 2)
        )


def synchronized_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    """Associate each source sample at most once within the configured tolerance."""
    tolerance = data["request"].sync_tolerance_ms / 1000
    rows: list[dict[str, Any]] = []
    for session in data["sessions"]:
        session_id = session["id"]
        electrical = sorted(
            (point for point in data["electrical"] if point["session_id"] == session_id),
            key=lambda point: point["timestamp"],
        )
        temperatures = sorted(
            (point for point in data["temperatures"] if point["session_id"] == session_id),
            key=lambda point: point["timestamp"],
        )
        electrical_times = [point["timestamp"] for point in electrical]
        used_electrical: set[int] = set()
        for temperature in temperatures:
            nearest = None
            nearest_index = None
            index = bisect_left(electrical_times, temperature["timestamp"])
            candidate_indexes = [
                candidate_index
                for candidate_index in range(max(0, index - 2), min(len(electrical), index + 2))
                if candidate_index not in used_electrical
            ]
            candidates = [
                (candidate_index, electrical[candidate_index])
                for candidate_index in candidate_indexes
            ]
            if candidates:
                nearest_index, candidate = min(
                    candidates,
                    key=lambda item: abs(
                        (item[1]["timestamp"] - temperature["timestamp"]).total_seconds()
                    ),
                )
                if (
                    abs((candidate["timestamp"] - temperature["timestamp"]).total_seconds())
                    <= tolerance
                ):
                    nearest = candidate
                    used_electrical.add(nearest_index)
            rows.append(
                {
                    "session_id": session_id,
                    "timestamp": temperature["timestamp"],
                    "temperature_sample_timestamp": temperature["timestamp"],
                    "electrical_sample_timestamp": nearest["timestamp"] if nearest else None,
                    "channels": temperature["channels"],
                    "ambient_temperature_c": temperature["ambient_temperature_c"],
                    "active_power_w": nearest["active_power_w"] if nearest else None,
                    "voltage_v": nearest["voltage_v"] if nearest else None,
                    "current_a": nearest["current_a"] if nearest else None,
                    "apparent_power_va": nearest["apparent_power_va"] if nearest else None,
                    "reactive_power_var": nearest["reactive_power_var"] if nearest else None,
                    "power_factor": nearest["power_factor"] if nearest else None,
                    "voltage_frequency_hz": nearest["voltage_frequency_hz"] if nearest else None,
                }
            )
        for electrical_index, point in enumerate(electrical):
            if electrical_index not in used_electrical:
                rows.append(
                    {
                        "session_id": session_id,
                        "timestamp": point["timestamp"],
                        "temperature_sample_timestamp": None,
                        "electrical_sample_timestamp": point["timestamp"],
                        "channels": {},
                        "ambient_temperature_c": None,
                        **{
                            key: point[key]
                            for key in (
                                "active_power_w",
                                "voltage_v",
                                "current_a",
                                "apparent_power_va",
                                "reactive_power_var",
                                "power_factor",
                                "voltage_frequency_hz",
                            )
                        },
                    }
                )
    rows.sort(key=lambda row: (row["session_id"], row["timestamp"]))
    return rows


def comparative_rows(data: dict[str, Any], request: PeriodReportRequest) -> list[dict[str, Any]]:
    """Build independent chart rows; no source sample is interpolated or duplicated."""
    rows: list[dict[str, Any]] = []
    for session in data["sessions"]:
        session_id = session["id"]
        origin = datetime.fromisoformat(session["started_at"])
        electrical = sorted(
            (point for point in data["electrical"] if point["session_id"] == session_id),
            key=lambda point: point["timestamp"],
        )
        temperatures = sorted(
            (point for point in data["temperatures"] if point["session_id"] == session_id),
            key=lambda point: point["timestamp"],
        )
        merged: dict[datetime | float, dict[str, Any]] = {}

        for point in electrical:
            axis_value: datetime | float = (
                (point["timestamp"] - origin).total_seconds()
                if request.time_axis_mode == "synchronized"
                else point["timestamp"]
            )
            row = merged.setdefault(
                axis_value,
                {"session_id": session_id, "axis_value": axis_value, "channels": {}},
            )
            row.update(
                {
                    "electrical_timestamp": point["timestamp"],
                    "active_power_w": point["active_power_w"],
                }
            )

        for point in temperatures:
            axis_value = (
                (point["timestamp"] - origin).total_seconds()
                if request.time_axis_mode == "synchronized"
                else point["timestamp"]
            )
            row = merged.setdefault(
                axis_value,
                {"session_id": session_id, "axis_value": axis_value, "channels": {}},
            )
            row["thermal_timestamp"] = point["timestamp"]
            row["channels"] = point["channels"]

        rows.extend(merged[key] for key in sorted(merged))
    return rows


def _append_synchronized(
    sheet: Any,
    rows: list[dict[str, Any]],
    channels: list[int],
    zone: ZoneInfo,
) -> None:
    sheet.append(
        [
            "Sessão",
            "Timestamp",
            "Potência ativa (W)",
            *[f"T{channel} (°C)" for channel in channels],
            "Tensão (V)",
            "Corrente (A)",
            "Potência aparente (VA)",
            "Potência reativa (var)",
            "Fator de potência",
            "Frequência (Hz)",
        ]
    )
    for row in rows:
        sheet.append(
            [
                row["session_id"],
                _local_excel(row["timestamp"], zone),
                row.get("active_power_w"),
                *[row["channels"].get(channel) for channel in channels],
                row.get("voltage_v"),
                row.get("current_a"),
                row.get("apparent_power_va"),
                row.get("reactive_power_var"),
                row.get("power_factor"),
                row.get("voltage_frequency_hz"),
            ]
        )
    _header(sheet, 1, sheet.max_column)
    for cell in sheet["B"][1:]:
        cell.number_format = "dd/mm/yyyy hh:mm:ss"
    _finish_sheet(sheet)


def _append_comparative(
    sheet: Any,
    rows: list[dict[str, Any]],
    channels: list[int],
    zone: ZoneInfo,
    mode: str,
) -> None:
    sheet.append(
        [
            "Sessão",
            "Tempo decorrido" if mode == "synchronized" else "Horário real",
            "Horário elétrico original",
            "Horário térmico original",
            "Potência ativa (W)",
            *[f"T{channel} (°C)" for channel in channels],
        ]
    )
    for row in rows:
        axis_value = row["axis_value"]
        if mode == "synchronized":
            axis_value = float(axis_value) / 86_400
        else:
            axis_value = _local_excel(axis_value, zone)
        sheet.append(
            [
                row["session_id"],
                axis_value,
                _local_excel(row["electrical_timestamp"], zone)
                if row.get("electrical_timestamp")
                else None,
                _local_excel(row["thermal_timestamp"], zone)
                if row.get("thermal_timestamp")
                else None,
                row.get("active_power_w"),
                *[row["channels"].get(channel) for channel in channels],
            ]
        )
    _header(sheet, 1, sheet.max_column)
    for cell in sheet["B"][1:]:
        cell.number_format = "[h]:mm:ss" if mode == "synchronized" else "dd/mm/yyyy hh:mm:ss"
    for column in ("C", "D"):
        for cell in sheet[column][1:]:
            cell.number_format = "dd/mm/yyyy hh:mm:ss.000"
    _finish_sheet(sheet)


def _add_curve_chart(sheet: Any, channels: list[int], anchor: str, mode: str) -> None:
    if sheet.max_row < 2:
        return
    # Each XY series contains only that source's actual observations, preserving gaps/Open.
    coordinates = sheet.parent.create_sheet("Coordenadas " + str(len(sheet.parent.worksheets)))
    coordinates.sheet_state = "hidden"
    temperature_chart = ScatterChart()
    temperature_chart.title = "Temperaturas e pot\u00eancia \u2014 in\u00edcio comum da sess\u00e3o"
    temperature_chart.x_axis.title = (
        "Tempo decorrido" if mode == "synchronized" else "Hor\u00e1rio real"
    )
    temperature_chart.x_axis.numFmt = "[h]:mm:ss" if mode == "synchronized" else "hh:mm:ss"
    temperature_chart.y_axis.title = "Temperatura (\u00b0C)"
    temperature_chart.height, temperature_chart.width = 11, 24
    power_chart = ScatterChart()
    thermal_values, power_values = [], []
    for index, channel in enumerate([None, *channels]):
        source_col, value_col = (3, 5) if channel is None else (4, 6 + channels.index(channel))
        col = index * 2 + 1
        label = "Pot\u00eancia (W)" if channel is None else f"T{channel}"
        coordinates.cell(1, col, "Tempo")
        coordinates.cell(1, col + 1, label)
        count = 1
        for row in sheet.iter_rows(min_row=2):
            if row[source_col - 1].value is None:
                continue
            count += 1
            coordinates.cell(count, col, row[1].value)
            value = row[value_col - 1].value
            coordinates.cell(count, col + 1, value)
            if isinstance(value, float | int):
                (power_values if channel is None else thermal_values).append(value)
        if count < 2:
            continue
        series = Series(
            Reference(coordinates, min_col=col + 1, min_row=2, max_row=count),
            Reference(coordinates, min_col=col, min_row=2, max_row=count),
            title=label,
        )
        color = (
            POWER_COLOR if channel is None else CHANNEL_COLORS[(channel - 1) % len(CHANNEL_COLORS)]
        )
        series.graphicalProperties.line.solidFill = color.lstrip("#")
        series.marker.symbol = "circle"
        series.marker.size = 2
        series.marker.graphicalProperties.solidFill = color.lstrip("#")
        (power_chart if channel is None else temperature_chart).series.append(series)
    temperature_chart.display_blanks = power_chart.display_blanks = "gap"
    temperature_chart.y_axis.scaling.min, temperature_chart.y_axis.scaling.max = padded_bounds(
        thermal_values
    )
    power_chart.y_axis.scaling.min, power_chart.y_axis.scaling.max = padded_bounds(power_values)
    power_chart.y_axis.title = "Pot\u00eancia (W)"
    power_chart.y_axis.axId = 200
    power_chart.y_axis.crosses = "max"
    power_chart.y_axis.majorGridlines = None
    temperature_chart += power_chart
    sheet.add_chart(temperature_chart, anchor)


def render_period_xlsx(data: dict[str, Any], request: PeriodReportRequest) -> bytes:
    zone = ZoneInfo(request.timezone)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo Executivo"
    channels = data["selected_channels"]
    channel_stats = [item for item in data["statistics"]["channels"] if item["count"]]
    general = data["statistics"]["general"]
    electrical = data["statistics"]["electrical"]
    temperature = data["statistics"]["temperature"]
    summary.merge_cells("A1:H2")
    summary["A1"] = "THERMOPOWER MONITOR\nRelatório Técnico de Ensaio Térmico e Elétrico"
    summary["A1"].font = Font(size=18, bold=True, color=WHITE)
    summary["A1"].fill = PatternFill("solid", fgColor=NAVY)
    summary["A1"].alignment = Alignment(vertical="center", wrap_text=True)
    summary["A4"] = request.title
    summary["A4"].font = Font(size=15, bold=True, color=NAVY)
    summary["A5"] = "Período analisado"
    summary["B5"] = _local_excel(request.start, zone)
    summary["C5"] = _local_excel(request.end, zone)
    summary["B5"].number_format = summary["C5"].number_format = "dd/mm/yyyy hh:mm:ss"
    kpis = [
        ("Período analisado", format_duration_pt(general["analyzed_period_seconds"])),
        (
            "Integridade dos dados",
            "Cobertura incompleta"
            if general.get("source_issue_count")
            else "Íntegra"
            if not general["gap_count"]
            else f"{general['gap_count']} lacuna(s)",
        ),
        ("Potência média (W)", electrical["active_power_w"]["mean"]),
        ("Potência máxima (W)", electrical["active_power_w"]["max"]),
        ("Energia (Wh)", electrical["energy_wh"]),
        ("Temperatura máxima (°C)", temperature["max"]),
        ("Canal crítico", temperature["critical_channel_label"]),
        (
            "ΔT máximo (°C)",
            temperature["maximum_delta_t"]["value_c"] if temperature["maximum_delta_t"] else None,
        ),
    ]
    for index, (label, value) in enumerate(kpis):
        row = 7 + index // 4 * 3
        column = 1 + (index % 4) * 2
        summary.cell(row, column, label)
        summary.cell(row + 1, column, value)
        summary.cell(row, column).font = Font(size=9, color="64748B", bold=True)
        summary.cell(row + 1, column).font = Font(size=14, color=NAVY, bold=True)
        summary.cell(row, column).fill = summary.cell(row + 1, column).fill = PatternFill(
            "solid", fgColor=PALE_BLUE
        )
    table_row = 14
    summary.append([])
    summary.cell(table_row, 1, "Canal")
    summary.cell(table_row, 2, "Identificação")
    summary.cell(table_row, 3, "Leituras")
    summary.cell(table_row, 4, "Média (°C)")
    summary.cell(table_row, 5, "Máx (°C)")
    summary.cell(table_row, 6, "Mín (°C)")
    summary.cell(table_row, 7, "ΔT (°C)")
    summary.cell(table_row, 8, "P95 (°C) — 95% abaixo")
    _header(summary, table_row, 8)
    for item in channel_stats:
        summary.append(
            [
                f"T{item['channel']}",
                item["friendly_name"],
                item["count"],
                item["mean"],
                item["max"],
                item["min"],
                item["range"],
                item["p95"],
            ]
        )
    summary.column_dimensions["A"].width = 22
    for column in "BCDEFGH":
        summary.column_dimensions[column].width = 18
    summary.freeze_panes = "A5"

    synchronized = synchronized_rows(data)
    comparative = comparative_rows(data, request)
    curves = workbook.create_sheet("Curvas do Ensaio")
    _append_comparative(curves, comparative, channels, zone, request.time_axis_mode)
    _add_curve_chart(curves, channels, "A4", request.time_axis_mode)

    stabilized = workbook.create_sheet("Análise Estabilizada")
    _append_synchronized(stabilized, synchronized, channels, zone)

    channel_sheet = workbook.create_sheet("Estatística por Canal")
    channel_sheet.append(
        [
            "Canal",
            "Nome",
            "Leituras",
            "Média",
            "Mínimo",
            "Máximo",
            "Mediana",
            "Desvio padrão",
            "P95 — 95% abaixo",
            "Amplitude",
            "Taxa °C/min",
            "Disponibilidade (%)",
        ]
    )
    for item in data["statistics"]["channels"]:
        channel_sheet.append(
            [
                f"T{item['channel']}",
                item["friendly_name"],
                item["count"],
                item["mean"],
                item["min"],
                item["max"],
                item["median"],
                item["standard_deviation"],
                item["p95"],
                item["range"],
                item["heating_rate_c_per_minute"],
                item["availability_percent"],
            ]
        )
    _header(channel_sheet, 1, channel_sheet.max_column)
    _finish_sheet(channel_sheet)

    electric_sheet = workbook.create_sheet("Grandezas Elétricas")
    electric_sheet.append(
        [
            "Grandeza",
            "Leituras",
            "Média",
            "Mínimo",
            "Máximo",
            "Mediana",
            "Desvio",
            "P95 — 95% abaixo",
        ]
    )
    electric_labels = {
        "active_power_w": "Potência ativa (W)",
        "voltage_v": "Tensão (V)",
        "current_a": "Corrente (A)",
        "apparent_power_va": "Potência aparente (VA)",
        "reactive_power_var": "Potência reativa (var)",
        "power_factor": "Fator de potência",
        "voltage_frequency_hz": "Frequência (Hz)",
    }
    for key, label in electric_labels.items():
        item = electrical[key]
        electric_sheet.append(
            [
                label,
                item["count"],
                item["mean"],
                item["min"],
                item["max"],
                item["median"],
                item["standard_deviation"],
                item["p95"],
            ]
        )
    electric_sheet.append(["Energia (Wh)", None, electrical["energy_wh"]])
    _header(electric_sheet, 1, electric_sheet.max_column)
    _finish_sheet(electric_sheet)

    real_electrical = workbook.create_sheet("Leituras Elétricas Reais")
    real_electrical.append(
        [
            "Sessão",
            "Timestamp original",
            "Tensão (V)",
            "Corrente (A)",
            "Potência ativa (W)",
            "Potência aparente (VA)",
            "Potência reativa (var)",
            "Fator de potência",
            "Frequência tensão (Hz)",
            "Frequência corrente (Hz)",
            "Qualidade",
            "Recebimento UTC",
            "Relógio do instrumento UTC",
        ]
    )
    for point in data["electrical"]:
        real_electrical.append(
            [
                point["session_id"],
                _local_excel(point["timestamp"], zone),
                point["voltage_v"],
                point["current_a"],
                point["active_power_w"],
                point["apparent_power_va"],
                point["reactive_power_var"],
                point["power_factor"],
                point["voltage_frequency_hz"],
                point["current_frequency_hz"],
                point["quality"],
                point.get("received_timestamp").isoformat()
                if point.get("received_timestamp")
                else None,
                point.get("device_timestamp").isoformat()
                if point.get("device_timestamp")
                else None,
            ]
        )
    _header(real_electrical, 1, real_electrical.max_column)
    for cell in real_electrical["B"][1:]:
        cell.number_format = "dd/mm/yyyy hh:mm:ss.000"
    _finish_sheet(real_electrical)

    real_thermal = workbook.create_sheet("Leituras Térmicas Reais")
    real_thermal.append(
        [
            "Sessão",
            "Timestamp original",
            "Ambiente (°C)",
            *[f"T{channel} (°C)" for channel in channels],
            "Qualidade",
            "Recebimento UTC",
            "Relógio do instrumento UTC",
        ]
    )
    for point in data["temperatures"]:
        real_thermal.append(
            [
                point["session_id"],
                _local_excel(point["timestamp"], zone),
                point["ambient_temperature_c"],
                *[point["channels"].get(channel) for channel in channels],
                point["quality"],
                point.get("received_timestamp").isoformat()
                if point.get("received_timestamp")
                else None,
                point.get("device_timestamp").isoformat()
                if point.get("device_timestamp")
                else None,
            ]
        )
    _header(real_thermal, 1, real_thermal.max_column)
    for cell in real_thermal["B"][1:]:
        cell.number_format = "dd/mm/yyyy hh:mm:ss.000"
    _finish_sheet(real_thermal)

    synchronized_sheet = workbook.create_sheet("Dados Sincronizados")
    _append_synchronized(synchronized_sheet, synchronized, channels, zone)

    metadata_sheet = workbook.create_sheet("Metadados")
    metadata_sheet.append(["Campo", "Valor"])
    metadata_sheet.append(["Versão", APPLICATION_VERSION])
    metadata_sheet.append(["Fuso horário", request.timezone])
    metadata_sheet.append(
        [
            "Eixo dos gráficos",
            "Início comum da sessão (apenas visual)"
            if request.time_axis_mode == "synchronized"
            else "Horário real",
        ]
    )
    metadata_sheet.append(["Início analisado", _local_excel(request.start, zone)])
    metadata_sheet.append(["Fim analisado", _local_excel(request.end, zone)])
    for session in data["sessions"]:
        metadata_sheet.append([f"Sessão #{session['id']}", session["name"]])
        metadata_sheet.append(["Operador", session["operator"] or "Operador não informado"])
        for key, value in session.get("metadata", {}).items():
            if value not in (None, ""):
                metadata_sheet.append(
                    [
                        key,
                        json.dumps(value, ensure_ascii=False)
                        if isinstance(value, dict | list)
                        else value,
                    ]
                )
        for device in session["devices"]:
            identity = " ".join(
                str(value) for value in (device.get("manufacturer"), device.get("model")) if value
            )
            metadata_sheet.append(
                [f"Fonte {device.get('role', 'combined')}", identity or device["name"]]
            )
            for key in ("serial_number", "port", "baud_rate", "cadence_ms"):
                if device.get(key) is not None:
                    metadata_sheet.append([f"  {key}", device[key]])
    _header(metadata_sheet, 1, 2)
    _finish_sheet(metadata_sheet)

    if electrical.get("peak"):
        peak = electrical["peak"]
        summary.append(
            [
                "Pico de potência (W)",
                peak["value_w"],
                "Tempo decorrido",
                format_duration_pt(peak["elapsed_seconds"]),
            ]
        )
    if temperature.get("critical_timestamp"):
        summary.append(
            [
                "Temperatura máxima",
                temperature["critical_value_c"],
                temperature["critical_channel_label"],
                temperature["critical_timestamp"],
            ]
        )

    # Reuse the same chart definition in the executive sheet while keeping its source data intact.
    if curves._charts:
        summary.add_chart(copy(curves._charts[0]), "J4")
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_period_csv(
    data: dict[str, Any], request: PeriodReportRequest, dataset: str = "synchronized"
) -> bytes:
    zone = ZoneInfo(request.timezone)
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, delimiter=";", lineterminator="\n")
    channels = data["selected_channels"]
    if dataset == "electrical":
        writer.writerow(
            [
                "sessao",
                "timestamp_original",
                "tensao_v",
                "corrente_a",
                "potencia_ativa_w",
                "potencia_aparente_va",
                "potencia_reativa_var",
                "fator_potencia",
                "frequencia_hz",
                "qualidade",
            ]
        )
        for point in data["electrical"]:
            writer.writerow(
                [
                    point["session_id"],
                    point["timestamp"].astimezone(zone).isoformat(),
                    point["voltage_v"],
                    point["current_a"],
                    point["active_power_w"],
                    point["apparent_power_va"],
                    point["reactive_power_var"],
                    point["power_factor"],
                    point["voltage_frequency_hz"],
                    point["quality"],
                ]
            )
    elif dataset == "thermal":
        writer.writerow(
            [
                "sessao",
                "timestamp_original",
                "temperatura_ambiente_c",
                *[f"temperatura_t{channel}_c" for channel in channels],
                "qualidade",
            ]
        )
        for point in data["temperatures"]:
            writer.writerow(
                [
                    point["session_id"],
                    point["timestamp"].astimezone(zone).isoformat(),
                    point["ambient_temperature_c"],
                    *[point["channels"].get(channel) for channel in channels],
                    point["quality"],
                ]
            )
    else:
        writer.writerow(
            [
                "sessao",
                "timestamp_grade",
                "timestamp_temperatura",
                "timestamp_eletrica",
                "potencia_ativa_w",
                *[f"temperatura_t{channel}_c" for channel in channels],
            ]
        )
        for row in synchronized_rows(data):
            writer.writerow(
                [
                    row["session_id"],
                    row["timestamp"].astimezone(zone).isoformat(),
                    row.get("temperature_sample_timestamp").astimezone(zone).isoformat()
                    if row.get("temperature_sample_timestamp")
                    else None,
                    row.get("electrical_sample_timestamp").astimezone(zone).isoformat()
                    if row.get("electrical_sample_timestamp")
                    else None,
                    row.get("active_power_w"),
                    *[row["channels"].get(channel) for channel in channels],
                ]
            )
    return ("\ufeff" + stream.getvalue()).encode("utf-8")
