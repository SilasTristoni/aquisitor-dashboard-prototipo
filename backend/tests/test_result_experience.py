import copy
import io
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from test_period_reports import _seed_period_data

from app.core.database import SessionLocal
from app.models.entities import (
    Measurement,
    MeasurementSession,
    SessionChannelConfiguration,
    TemperatureChannelValue,
    TemperatureMeasurement,
    TemperatureSample,
)
from app.schemas.contracts import PeriodReportRequest
from app.services.period_documents import _segments, render_period_chart, render_period_pdf
from app.services.period_reporting import PeriodReportDataService, period_statistics
from app.services.period_workbook import render_period_csv, render_period_xlsx


def seed_selected():
    start, end, ids = _seed_period_data()
    with SessionLocal() as db:
        for channel in [25, 29, 32]:
            db.add(
                SessionChannelConfiguration(
                    session_id=ids[0],
                    channel=channel,
                    name=f"Sensor {channel}",
                    enabled=True,
                    sensor_type="K",
                    unit="°C",
                    correction_offset=0,
                    color="#2563eb",
                    display_order=channel,
                )
            )
        samples = list(
            db.scalars(
                select(TemperatureSample)
                .where(TemperatureSample.session_id == ids[0])
                .order_by(TemperatureSample.received_timestamp)
            )
        )
        for sample in samples:
            for channel in [25, 29, 32]:
                sample.channels.append(
                    TemperatureChannelValue(
                        channel=channel,
                        temperature_c=80 if channel in [25, 29] else None,
                        original_value=80 if channel in [25, 29] else None,
                        original_unit="°C",
                        quality="good" if channel != 32 else "open",
                    )
                )
        db.commit()
    return start, end, ids


@pytest.mark.parametrize("channels", [[25, 29, 32], [29], None, [], [32]])
def test_selection_peaks_and_raw_preservation(client, channels):
    start, end, ids = seed_selected()
    request = PeriodReportRequest(
        start=start, end=end, session_ids=[ids[0]], channels=channels, include_open_channels=True
    )
    with SessionLocal() as db:
        service = PeriodReportDataService(db)
        data = service.collect(request)
        expected = channels if channels is not None else [1, 25, 29, 32]
        assert data["selected_channels"] == expected
        assert data["statistics"]["electrical"]["active_power_w"]["count"] == 3
        peak = data["statistics"]["temperature"]["peak"]
        if channels in ([], [32]):
            assert peak is None
        else:
            assert peak["timestamp"] == start.isoformat()
            assert peak["channel"] == (29 if channels == [29] else 25)
            assert peak["session_elapsed_seconds"] == peak["period_elapsed_seconds"] == 0
        before = copy.deepcopy(data["temperatures"])
        book = load_workbook(io.BytesIO(render_period_xlsx(data, request)))
        assert "Picos térmicos" in book.sheetnames
        if peak:
            assert book["Picos térmicos"]["A2"].value.startswith("Máxima global · ")
        assert render_period_csv(data, request)
        assert data["temperatures"] == before
        assert (
            db.scalar(select(TemperatureChannelValue).where(TemperatureChannelValue.channel == 1))
            is not None
        )
        preview = service.preview_data(data, request)
        assert preview["selected_channels"] == expected
        if channels == [29]:
            assert render_period_pdf(data, request).startswith(b"%PDF")
            assert render_period_chart(data, request).startswith(b"\x89PNG")
            thermal_only = request.model_copy(
                update={"include_power": False, "include_electrical_details": False}
            )
            thermal_data = service.collect(thermal_only)
            assert thermal_data["statistics"]["electrical"]["energy_wh"] is None
            assert render_period_pdf(thermal_data, thermal_only).startswith(b"%PDF")


def test_peak_period_origin_invalid_and_ties(client):
    start, end, ids = seed_selected()
    request = PeriodReportRequest(
        start=start + timedelta(seconds=5),
        end=end,
        session_ids=[ids[0]],
        channels=[25, 29, 32],
        include_open_channels=True,
    )
    with SessionLocal() as db:
        data = PeriodReportDataService(db).collect(request)
    peak = data["statistics"]["temperature"]["peak"]
    assert peak["session_elapsed_seconds"] == 10
    assert peak["period_elapsed_seconds"] == 5
    invalid = copy.deepcopy(data["temperatures"][0])
    invalid["channels"] = {25: float("inf"), 29: float("nan"), 32: None}
    data["temperatures"].append(invalid)
    invalid = copy.deepcopy(invalid)
    invalid.update(channels={25: 99999}, quality="invalid")
    data["temperatures"].append(invalid)
    assert period_statistics(data)["temperature"]["peak"] == peak


def test_gaps_derived_before_downsampling_zero_not_missing(client, monkeypatch):
    start, end, ids = seed_selected()
    request = PeriodReportRequest(start=start, end=end, session_ids=[ids[0]])
    with SessionLocal() as db:
        service = PeriodReportDataService(db)
        data = service.collect(request)
        points = data["electrical"]
        points[0]["active_power_w"] = 0
        data["sessions"][0]["devices"] = [{"role": "electrical", "cadence_ms": 1000}]
        # Enough regular intervals to distinguish a physical gap from a slower valid source.
        points[1]["timestamp"] = start + timedelta(seconds=1)
        points.insert(2, {**points[1], "timestamp": start + timedelta(seconds=2)})
        preview = service.preview_data(data, request)
        assert preview["series"][0]["gaps"][0] == {
            "source": "electrical",
            "after": (start + timedelta(seconds=2)).isoformat(),
            "before": (start + timedelta(seconds=120)).isoformat(),
        }
        assert preview["series"][0]["electrical"][0]["active_power_w"] == 0
        segments = _segments(
            [
                {"timestamp": start + timedelta(seconds=i), "v": v}
                for i, v in enumerate([0, None, 1])
            ],
            60,
            "v",
        )
        assert [[p["v"] for p in part] for part in segments] == [[0], [1]]


def test_legacy_invalid_channel_excluded_without_rewriting_stored_sample(client):
    start, end, ids = _seed_period_data()
    with SessionLocal() as db:
        existing = db.get(MeasurementSession, ids[0])
        session = MeasurementSession(
            device_id=existing.device_id,
            user_id=existing.user_id,
            name="Legacy quality",
            started_at=start,
            ended_at=end,
            status="finished",
        )
        db.add(session)
        db.flush()
        measurement = Measurement(
            session_id=session.id,
            timestamp=start + timedelta(seconds=10),
            power_w=0,
            raw_power=0,
            raw_power_unit="W",
            quality="good",
        )
        measurement.temperatures.extend(
            [
                TemperatureMeasurement(channel=25, temperature_c=80, quality="good"),
                TemperatureMeasurement(channel=29, temperature_c=999, quality="invalid"),
                TemperatureMeasurement(channel=32, temperature_c=None, quality="open"),
            ]
        )
        db.add(measurement)
        db.commit()
        data = PeriodReportDataService(db).collect(
            PeriodReportRequest(
                start=start,
                end=end,
                session_ids=[session.id],
                channels=[25, 29, 32],
                include_open_channels=True,
            )
        )
        assert data["statistics"]["temperature"]["peak"]["channel"] == 25
        assert data["statistics"]["electrical"]["active_power_w"]["mean"] == 0
        assert data["temperatures"][0]["channels"] == {25: 80, 29: None, 32: None}
        db.refresh(measurement)
        assert (
            next(item.temperature_c for item in measurement.temperatures if item.channel == 29)
            == 999
        )
