import io
from datetime import UTC, datetime, timedelta
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import select

from app.core.database import SessionLocal
from app.models.entities import (
    Device,
    ElectricalSample,
    Measurement,
    MeasurementSession,
    Report,
    SessionChannelConfiguration,
    SessionDevice,
    TemperatureChannelValue,
    TemperatureMeasurement,
    TemperatureSample,
    User,
)
from app.schemas.contracts import PeriodReportRequest
from app.services.period_reporting import PeriodReportDataService, downsample_time_buckets
from app.services.usb_discovery import usb_discovery_service


def _seed_period_data() -> tuple[datetime, datetime, list[int]]:
    start = datetime(2026, 1, 15, 13, 0, tzinfo=UTC)
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.email == "admin@demo.thermopower.com"))
        device = db.scalar(select(Device).where(Device.name == "Aquisitor simulado"))
        sessions = [
            MeasurementSession(
                device_id=device.id,
                user_id=user.id,
                name="Período A",
                started_at=start,
                ended_at=start + timedelta(minutes=3),
                status="finished",
            ),
            MeasurementSession(
                device_id=device.id,
                user_id=user.id,
                name="Período B concorrente",
                started_at=start + timedelta(seconds=5),
                ended_at=start + timedelta(minutes=2),
                status="finished",
            ),
        ]
        db.add_all(sessions)
        db.flush()
        for session in sessions:
            db.add(
                SessionChannelConfiguration(
                    session_id=session.id,
                    channel=1,
                    name=f"Entrada {session.id}",
                    enabled=True,
                    sensor_type="K",
                    unit="°C",
                    correction_offset=0,
                    color="#2563EB",
                    display_order=1,
                )
            )

        electrical_rows = [
            (sessions[0], 0, 100.0, True),
            (sessions[0], 10, 200.0, True),
            (sessions[0], 120, 300.0, False),
            (sessions[1], 5, 400.0, True),
            (sessions[1], 15, 400.0, True),
        ]
        for session, seconds, power, has_device_timestamp in electrical_rows:
            timestamp = start + timedelta(seconds=seconds)
            db.add(
                ElectricalSample(
                    session_id=session.id,
                    device_id=device.id,
                    device_timestamp=timestamp if has_device_timestamp else None,
                    received_timestamp=timestamp,
                    active_power_w=power,
                    voltage_v=220,
                    current_a=power / 220,
                    power_factor=0.95,
                    quality="good",
                    source="test",
                    original_values={"active_power": power},
                    original_units={"active_power": "W"},
                    raw_payload={},
                )
            )
        for session, seconds, temperature in [
            (sessions[0], 0, 30.0),
            (sessions[0], 10, 32.0),
            (sessions[1], 5, 40.0),
            (sessions[1], 15, 42.0),
        ]:
            timestamp = start + timedelta(seconds=seconds)
            sample = TemperatureSample(
                session_id=session.id,
                device_id=device.id,
                device_timestamp=timestamp,
                received_timestamp=timestamp,
                quality="good",
                source="test",
                raw_payload={},
            )
            sample.channels.append(
                TemperatureChannelValue(
                    channel=1,
                    temperature_c=temperature,
                    original_value=temperature,
                    original_unit="°C",
                    quality="good",
                )
            )
            db.add(sample)
        db.commit()
        return start, start + timedelta(minutes=3), [session.id for session in sessions]


def _payload(start: datetime, end: datetime) -> dict:
    return {
        "start": start.isoformat(),
        "end": end.isoformat(),
        "timezone": "America/Sao_Paulo",
        "title": "Ensaio período A/B",
        "channels": [1],
        "table_max_rows": 20,
        "dpi": 96,
    }


def _seed_client_preview_data() -> tuple[datetime, datetime, int, int]:
    start = datetime(2026, 9, 3, 13, 0, tzinfo=UTC)
    timestamps = [start + timedelta(seconds=30 * index) for index in range(8)]
    thermal_offset = timedelta(milliseconds=350)
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.email == "admin@demo.thermopower.com"))
        thermal = Device(
            name="Termômetro multiponto",
            manufacturer="Applent",
            model="AT4532",
            serial_number="AT-CLIENT-PREVIEW",
            connection_type="serial",
            port="COM5",
            baud_rate=19200,
            protocol="at4532_serial",
            metadata_json={"expected_interval_ms": 1000},
        )
        electrical = Device(
            name="Medidor de potência",
            manufacturer="GW Instek",
            model="GPM-8213",
            serial_number="GPM-CLIENT-PREVIEW",
            connection_type="serial",
            port="COM4",
            baud_rate=9600,
            protocol="gpm8213_serial",
            metadata_json={"expected_interval_ms": 1000},
        )
        db.add_all([thermal, electrical])
        db.flush()
        session = MeasurementSession(
            device_id=thermal.id,
            user_id=user.id,
            name="Ensaio combinado Britânia",
            description="Validação térmica e elétrica",
            started_at=start,
            ended_at=timestamps[-1] + thermal_offset,
            status="finished",
            metadata_json={
                "product": "Forno elétrico",
                "model": "BFE50",
                "sample": "Amostra 04",
                "code": "ENG-2026-004",
                "nominal_voltage": "220 V",
                "responsible": "Engenharia de Produto",
            },
        )
        db.add(session)
        db.flush()
        db.add_all(
            [
                SessionDevice(session_id=session.id, device_id=thermal.id, role="temperature"),
                SessionDevice(session_id=session.id, device_id=electrical.id, role="electrical"),
            ]
        )
        for order, (channel, name) in enumerate(
            [(25, "Saída de ar"), (26, "Carcaça superior"), (27, "T27")], start=1
        ):
            db.add(
                SessionChannelConfiguration(
                    session_id=session.id,
                    channel=channel,
                    name=name,
                    enabled=True,
                    sensor_type="K",
                    unit="°C",
                    correction_offset=0,
                    color="#2563EB",
                    display_order=order,
                )
            )
        for index, timestamp in enumerate(timestamps):
            power = 780.0 if index % 2 == 0 else 180.0
            db.add(
                ElectricalSample(
                    session_id=session.id,
                    device_id=electrical.id,
                    device_timestamp=timestamp,
                    received_timestamp=timestamp + timedelta(milliseconds=80),
                    voltage_v=220.0,
                    current_a=power / 220,
                    active_power_w=power,
                    apparent_power_va=power / 0.95,
                    reactive_power_var=120.0,
                    power_factor=0.95,
                    voltage_frequency_hz=60.0,
                    current_frequency_hz=60.0,
                    original_values={"active_power": power},
                    original_units={"active_power": "W"},
                    quality="good",
                    source="test",
                    raw_payload={},
                )
            )
            thermal_timestamp = timestamp + thermal_offset
            thermal_sample = TemperatureSample(
                session_id=session.id,
                device_id=thermal.id,
                device_timestamp=thermal_timestamp,
                received_timestamp=thermal_timestamp + timedelta(milliseconds=40),
                ambient_temperature_c=23.0,
                quality="good",
                source="test",
                raw_payload={},
            )
            for channel, value in (
                (25, 60.0 + index * 0.1),
                (26, 50.0 + index * 0.05),
                (27, None),
            ):
                thermal_sample.channels.append(
                    TemperatureChannelValue(
                        channel=channel,
                        temperature_c=value,
                        original_value=value,
                        original_unit="°C",
                        quality="good" if value is not None else "open",
                    )
                )
            db.add(thermal_sample)
        db.commit()
        return start, timestamps[-1] + thermal_offset, session.id, len(timestamps)


def test_period_contract_normalizes_naive_local_time_and_validates_metrics():
    request = PeriodReportRequest(
        start=datetime(2026, 1, 15, 10, 0),
        end=datetime(2026, 1, 15, 11, 0),
        channels=[2, 1, 2],
    )
    assert request.start == datetime(2026, 1, 15, 13, 0, tzinfo=UTC)
    assert request.channels == [1, 2]
    with pytest.raises(ValueError, match="posterior"):
        PeriodReportRequest(start=request.end, end=request.start)
    with pytest.raises(ValueError, match="ao menos uma métrica"):
        PeriodReportRequest(
            start=request.start,
            end=request.end,
            include_power=False,
            include_temperatures=False,
            include_electrical_details=False,
        )


def test_period_statistics_do_not_integrate_between_sessions(client: TestClient):
    start, end, session_ids = _seed_period_data()
    request = PeriodReportRequest(**_payload(start, end))
    with SessionLocal() as db:
        data = PeriodReportDataService(db).collect(request)
    assert [session["id"] for session in data["sessions"]] == session_ids
    assert data["statistics"]["general"]["session_count"] == 2
    assert data["statistics"]["general"]["gap_count"] >= 1
    assert data["statistics"]["general"]["timestamp_fallback_count"] == 1
    assert data["statistics"]["electrical"]["excluded_energy_intervals"] == 1
    assert data["statistics"]["electrical"]["energy_wh"] == pytest.approx(1.527777, rel=1e-5)
    assert {row["session_id"] for row in data["table_rows"]} == set(session_ids)


def test_period_endpoints_render_and_audit_files(client: TestClient, auth_headers: dict[str, str]):
    start, end, _ = _seed_period_data()
    payload = {**_payload(start, end), "channels": list(range(1, 10))}
    preview = client.post("/api/v1/reports/period/preview", headers=auth_headers, json=payload)
    assert preview.status_code == 200, preview.text
    assert len(preview.json()["series"]) == 2
    for endpoint, media, signature in [
        ("pdf", "application/pdf", b"%PDF"),
        ("chart.png", "image/png", b"\x89PNG"),
        ("chart.jpeg", "image/jpeg", b"\xff\xd8"),
    ]:
        response = client.post(
            f"/api/v1/reports/period/{endpoint}", headers=auth_headers, json=payload
        )
        assert response.status_code == 200, response.text
        assert media in response.headers["content-type"]
        assert response.content.startswith(signature)
        assert "attachment" in response.headers["content-disposition"]
        if endpoint == "pdf":
            assert b"/Subtype /Image" in response.content
            assert response.content.count(b"/Type /Page") >= 3
    history = client.get("/api/v1/reports", headers=auth_headers).json()
    period_reports = [row for row in history if row["scope_type"] == "period"]
    assert len(period_reports) == 3
    assert all(row["status"] == "completed" for row in period_reports)


def test_client_preview_auto_selects_active_channels_and_updates_identification(
    client: TestClient, auth_headers: dict[str, str]
):
    start, end, session_id, _ = _seed_client_preview_data()
    payload = {
        **_payload(start, end),
        "session_ids": [session_id],
        "channels": None,
    }
    preview = client.post(
        "/api/v1/reports/period/preview", headers=auth_headers, json=payload
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["selected_channels"] == [25, 26]
    assert body["channel_labels"] == {
        "25": "T25 — Saída de ar",
        "26": "T26 — Carcaça superior",
    }
    assert body["statistics"]["temperature"]["critical_channel"] == 25
    assert body["statistics"]["temperature"]["maximum_delta_t"]["value_c"] == pytest.approx(
        10.35
    )
    assert body["statistics"]["electrical"]["energy_wh"] > 0
    assert body["statistics"]["electrical"]["cycles"]["cycle_count"] == 4
    assert body["period"]["time_axis_mode"] == "synchronized"
    assert len(body["series"][0]["electrical"]) == len(body["series"][0]["temperatures"])
    first_electrical = body["series"][0]["electrical"][0]
    first_thermal = body["series"][0]["temperatures"][0]
    assert datetime.fromisoformat(first_thermal["timestamp"]) - datetime.fromisoformat(
        first_electrical["timestamp"]
    ) == timedelta(milliseconds=350)
    assert first_electrical["elapsed_seconds"] == 0
    assert first_thermal["elapsed_seconds"] == pytest.approx(0.35)
    assert body["series"][0]["session_started_at"] == start.isoformat()

    real_time = client.post(
        "/api/v1/reports/period/preview",
        headers=auth_headers,
        json={**payload, "time_axis_mode": "real"},
    )
    assert real_time.status_code == 200, real_time.text
    assert real_time.json()["period"]["time_axis_mode"] == "real"
    assert real_time.json()["statistics"] == body["statistics"]

    with_open = client.post(
        "/api/v1/reports/period/preview",
        headers=auth_headers,
        json={**payload, "include_open_channels": True},
    )
    assert with_open.status_code == 200, with_open.text
    assert with_open.json()["selected_channels"] == [25, 26, 27]
    open_stats = next(
        item for item in with_open.json()["statistics"]["channels"] if item["channel"] == 27
    )
    assert open_stats["count"] == 0
    assert open_stats["label"] == "T27"

    updated = client.patch(
        f"/api/v1/sessions/{session_id}",
        headers=auth_headers,
        json={
            "metadata": {
                "product": "Air fryer",
                "model": "BFR51",
                "responsible": "Laboratório Britânia",
            },
            "channel_names": {"25": "Saída traseira"},
        },
    )
    assert updated.status_code == 200, updated.text
    detail = client.get(f"/api/v1/sessions/{session_id}", headers=auth_headers)
    assert detail.status_code == 200
    assert detail.json()["metadata"] == {
        "product": "Air fryer",
        "model": "BFR51",
        "responsible": "Laboratório Britânia",
    }
    assert next(row for row in detail.json()["channels"] if row["channel"] == 25)["name"] == (
        "Saída traseira"
    )


def test_client_preview_outputs_keep_raw_streams_and_professional_workbook(
    client: TestClient, auth_headers: dict[str, str]
):
    start, end, session_id, sample_count = _seed_client_preview_data()
    payload = {
        **_payload(start, end),
        "session_ids": [session_id],
        "channels": None,
        "title": "Resumo executivo — Britânia",
    }
    spreadsheet = client.post(
        "/api/v1/reports/period/xlsx", headers=auth_headers, json=payload
    )
    assert spreadsheet.status_code == 200, spreadsheet.text
    workbook = load_workbook(io.BytesIO(spreadsheet.content), data_only=True)
    assert workbook.sheetnames == [
        "Resumo Executivo",
        "Curvas do Ensaio",
        "Análise Estabilizada",
        "Estatística por Canal",
        "Grandezas Elétricas",
        "Leituras Elétricas Reais",
        "Leituras Térmicas Reais",
        "Dados Sincronizados",
        "Metadados",
    ]
    assert workbook["Leituras Elétricas Reais"].max_row == sample_count + 1
    assert workbook["Leituras Térmicas Reais"].max_row == sample_count + 1
    assert workbook["Dados Sincronizados"].max_row == sample_count + 1
    assert isinstance(workbook["Leituras Elétricas Reais"]["E2"].value, int | float)
    assert isinstance(workbook["Leituras Térmicas Reais"]["D2"].value, int | float)
    assert len(workbook["Curvas do Ensaio"]._charts) == 1
    assert len(workbook["Resumo Executivo"]._charts) == 1
    curves = workbook["Curvas do Ensaio"]
    assert curves["B1"].value == "Tempo decorrido"
    assert curves["B2"].value == timedelta(milliseconds=350)
    assert curves["C2"].value != curves["D2"].value
    assert curves["E2"].value is not None
    assert curves["F2"].value is not None

    thermal_csv = client.post(
        "/api/v1/reports/period/csv?dataset=thermal", headers=auth_headers, json=payload
    )
    assert thermal_csv.status_code == 200
    assert thermal_csv.content.startswith(b"\xef\xbb\xbf")
    assert len(thermal_csv.content.decode("utf-8-sig").splitlines()) == sample_count + 1

    executive_png = client.post(
        "/api/v1/reports/period/executive.png", headers=auth_headers, json=payload
    )
    assert executive_png.status_code == 200, executive_png.text
    assert Image.open(io.BytesIO(executive_png.content)).size == (1200, 650)
    assert len(executive_png.content) > 40_000

    executive_pdf = client.post(
        "/api/v1/reports/period/executive.pdf", headers=auth_headers, json=payload
    )
    assert executive_pdf.status_code == 200, executive_pdf.text
    assert executive_pdf.content.startswith(b"%PDF")
    assert executive_pdf.content.count(b"/Type /Page") == 2


def test_period_endpoint_reports_clear_empty_data_error(
    client: TestClient, auth_headers: dict[str, str]
):
    response = client.post(
        "/api/v1/reports/period/preview",
        headers=auth_headers,
        json=_payload(datetime(2035, 1, 1, tzinfo=UTC), datetime(2035, 1, 2, tzinfo=UTC)),
    )
    assert response.status_code == 422
    assert "Nenhuma sessão" in response.json()["error"]["message"]


def test_downsampling_preserves_global_extremes():
    start = datetime(2026, 1, 1, tzinfo=UTC)
    points = [
        {"timestamp": start + timedelta(seconds=index), "active_power_w": float(index)}
        for index in range(1000)
    ]
    points[501]["active_power_w"] = -500
    reduced = downsample_time_buckets(points, ["active_power_w"], 80)
    values = [point["active_power_w"] for point in reduced]
    assert len(reduced) <= 80
    assert min(values) == -500
    assert max(values) == 999
    assert reduced[0]["timestamp"] == points[0]["timestamp"]
    assert reduced[-1]["timestamp"] == points[-1]["timestamp"]


def test_period_service_falls_back_to_unmigrated_legacy_rows(client: TestClient):
    start = datetime(2025, 4, 1, 12, 0, tzinfo=UTC)
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.email == "admin@demo.thermopower.com"))
        device = db.scalar(select(Device).where(Device.name == "Aquisitor simulado"))
        session = MeasurementSession(
            device_id=device.id,
            user_id=user.id,
            name="Somente legado",
            started_at=start,
            ended_at=start + timedelta(minutes=1),
            status="finished",
        )
        db.add(session)
        db.flush()
        measurement = Measurement(
            session_id=session.id,
            timestamp=start + timedelta(seconds=10),
            power_w=123.0,
            raw_power=123.0,
            raw_power_unit="W",
            quality="good",
        )
        measurement.temperatures.append(
            TemperatureMeasurement(channel=1, temperature_c=31.5, quality="good")
        )
        db.add(measurement)
        db.commit()
        request = PeriodReportRequest(
            start=start,
            end=start + timedelta(minutes=1),
            session_ids=[session.id],
            channels=[1],
        )
        data = PeriodReportDataService(db).collect(request)
    assert data["electrical"][0]["source"] == "legacy"
    assert data["temperatures"][0]["channels"] == {1: 31.5}


def test_usb_discovery_association_and_conservative_suggestion(
    client: TestClient, auth_headers: dict[str, str], monkeypatch: pytest.MonkeyPatch
):
    port = SimpleNamespace(
        device="COM9",
        description="USB AT4532 bridge",
        manufacturer="Laboratório",
        product="AT4532",
        serial_number="USB-TEST-9",
        vid=0x1234,
        pid=0x5678,
        hwid="USB VID:PID=1234:5678",
        location="1-4",
    )

    class FakeSerial:
        is_open = True

        def close(self):
            self.is_open = False

    busy_port = SimpleNamespace(
        device="COM10",
        description="Conversor USB serial",
        manufacturer=None,
        product=None,
        serial_number=None,
        vid=None,
        pid=None,
        hwid="USB UNKNOWN",
        location=None,
    )

    def serial_factory(**kwargs):
        if kwargs["port"] == "COM10":
            from serial import SerialException

            raise SerialException("Access denied: port is busy")
        return FakeSerial()

    monkeypatch.setattr(usb_discovery_service, "port_provider", lambda: [port, busy_port])
    monkeypatch.setattr(usb_discovery_service, "serial_factory", serial_factory)
    discovery = client.get("/api/v1/hardware/discovery", headers=auth_headers)
    assert discovery.status_code == 200
    item = next(row for row in discovery.json() if row["port"] == "COM9")
    assert item["status"] == "available"
    assert item["identification_status"] == "possible_at4532"
    assert item["confidence"] == "medium"
    assert item["association_status"] == "unassociated"
    busy = next(row for row in discovery.json() if row["port"] == "COM10")
    assert busy["status"] == "port_busy"
    assert busy["manufacturer"] is None
    device_id = client.get("/api/v1/devices", headers=auth_headers).json()[0]["id"]
    associated = client.post(
        "/api/v1/hardware/discovery/associate",
        headers=auth_headers,
        json={"port": "COM9", "device_id": device_id},
    )
    assert associated.status_code == 200
    assert associated.json()["metadata"]["usb"]["serial_number"] == "USB-TEST-9"
    second = next(
        row
        for row in client.get("/api/v1/hardware/discovery", headers=auth_headers).json()
        if row["port"] == "COM9"
    )
    assert second["association"]["matched_by"] == "serial_number"
    disposable = client.post(
        "/api/v1/devices", headers=auth_headers, json={"name": "Equipamento descartável"}
    )
    assert disposable.status_code == 201
    removed = client.delete(f"/api/v1/devices/{disposable.json()['id']}", headers=auth_headers)
    assert removed.status_code == 204
    active_ids = {row["id"] for row in client.get("/api/v1/devices", headers=auth_headers).json()}
    assert disposable.json()["id"] not in active_ids
    with SessionLocal() as db:
        report = db.scalar(select(Report).where(Report.scope_type == "period"))
        assert report is None
