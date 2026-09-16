"""Deterministic receipt-clock regression through the real service and HTTP exports."""

import asyncio
import io
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.adapters.base import DeviceReading, DeviceStatus
from app.core.database import SessionLocal
from app.models.entities import Device, ElectricalSample, MeasurementSession, TemperatureSample
from app.services.acquisition import AcquisitionService, DeviceRuntime, acquisition_service
from app.services.websocket import websocket_hub


class ClockedStream:
    expected_interval_seconds = 1.0

    def __init__(self, device_id, role, count):
        self.device_id, self.role, self.count = device_id, role, count
        self.connected = False
        self.last_message_at = None

    async def connect(self):
        self.connected = True

    async def disconnect(self):
        self.connected = False

    async def stop_reading(self):
        self.connected = False

    async def get_status(self):
        return DeviceStatus(
            state="reading",
            connected=self.connected,
            reading=self.connected,
            last_message_at=self.last_message_at,
        )

    async def start_reading(self):
        await asyncio.sleep(0.02)
        origin = datetime.now(UTC)
        for index in range(self.count):
            received = origin + timedelta(seconds=index)
            # A repeated, unsynchronized instrument clock must never collapse samples.
            instrument = origin - timedelta(minutes=5) + timedelta(seconds=index // 6 * 6)
            self.last_message_at = received
            thermal = self.role == "temperature"
            yield DeviceReading(
                timestamp=instrument if thermal else received,
                received_timestamp=received,
                device_timestamp=instrument if thermal else None,
                raw_power=None if thermal else 180 - index,
                raw_power_unit="W",
                power_w=None if thermal else 180 - index,
                temperatures_c=[None] * 24 + [22 + index / 2 + c / 10 for c in range(8)]
                if thermal
                else [],
            )
            while not acquisition_service.runtimes[self.device_id].session_id:
                await asyncio.sleep(0.001)
            await asyncio.sleep(0)
        while self.connected:
            await asyncio.sleep(0.01)


@pytest.mark.parametrize("count", [100, 120, 300])
def test_fresh_pair_to_database_api_websocket_and_reports(client, auth_headers, monkeypatch, count):
    with SessionLocal() as db:
        devices = [
            Device(
                name="Pipeline power",
                protocol="gpm8213_serial",
                connection_type="serial",
                port="PIPE1",
                baud_rate=9600,
            ),
            Device(
                name="Pipeline thermal",
                protocol="at4532_serial",
                connection_type="serial",
                port="PIPE2",
                baud_rate=19200,
            ),
        ]
        db.add_all(devices)
        db.commit()
        power_id, thermal_id = [d.id for d in devices]
    monkeypatch.setattr(
        acquisition_service,
        "_adapter_for",
        lambda d: ClockedStream(d.id, "electrical" if d.id == power_id else "temperature", count),
    )
    monkeypatch.setattr(websocket_hub, "queue_size", count * 4)
    token = auth_headers["Authorization"].split()[1]
    with client.websocket_connect(f"/api/v1/ws?token={token}") as ws:
        started = client.post(
            "/api/v1/sessions",
            headers=auth_headers,
            json={
                "name": "Deterministic 1 Hz",
                "electrical_device_id": power_id,
                "temperature_device_id": thermal_id,
            },
        )
        assert started.status_code == 201, started.text
        session_id = started.json()["id"]
        origin = datetime.fromisoformat(started.json()["started_at"])
        assert origin.tzinfo is not None
        messages = []
        while len(messages) < count * 2:
            event = ws.receive_json()
            assert event["type"] != "heartbeat", f"Only {len(messages)} session samples delivered"
            if (
                event["type"] == "measurement.created"
                and event["payload"]["session_id"] == session_id
            ):
                messages.append(event["payload"])
        assert len({(r["device_id"], r["received_timestamp"]) for r in messages}) == 2 * count
        for device_id in [power_id, thermal_id]:
            rows = [r for r in messages if r["device_id"] == device_id]
            assert len(rows) == count
            assert 0 <= (datetime.fromisoformat(rows[0]["timestamp"]) - origin).total_seconds() < 1
        assert (
            client.post(f"/api/v1/sessions/{session_id}/finish", headers=auth_headers).status_code
            == 200
        )
    # The fixture advances measurement time, not wall-clock execution time.
    with SessionLocal() as db:
        session = db.get(MeasurementSession, session_id)
        session.ended_at = origin + timedelta(seconds=count)
        db.commit()
        for model in [ElectricalSample, TemperatureSample]:
            assert (
                db.scalar(
                    select(func.count()).select_from(model).where(model.session_id == session_id)
                )
                == count
            )
        thermal = list(
            db.scalars(select(TemperatureSample).where(TemperatureSample.session_id == session_id))
        )
        assert len({r.sequence for r in thermal}) == count
        assert len({r.device_timestamp for r in thermal}) < count
    request = {
        "start": origin.isoformat(),
        "end": (origin + timedelta(seconds=count)).isoformat(),
        "session_ids": [session_id],
        "use_device_timestamp": True,
    }
    preview = client.post("/api/v1/reports/period/preview", headers=auth_headers, json=request)
    assert preview.status_code == 200, preview.text
    data = preview.json()
    assert data["statistics"]["general"]["electrical_sample_count"] == count
    assert data["statistics"]["general"]["temperature_sample_count"] == count
    assert data["statistics"]["electrical"]["peak"]["value_w"] == 180
    assert data["selected_channels"] == list(range(25, 33))
    for stream in ["electrical", "temperatures"]:
        assert len(data["series"][0][stream]) == count
    for extension in ["xlsx", "pdf", "png", "jpeg"]:
        response = client.get(
            f"/api/v1/reports/sessions/{session_id}.{extension}", headers=auth_headers
        )
        assert response.status_code == 200, response.text[:200]
        assert len(response.content) > 1000
        if extension == "xlsx":
            book = load_workbook(io.BytesIO(response.content), data_only=True)
            assert book["Leituras Elétricas Reais"].max_row == count + 1
            assert book["Leituras Térmicas Reais"].max_row == count + 1
            assert book["Leituras Térmicas Reais"].cell(2, 4).value == 22


@pytest.mark.asyncio
async def test_common_start_rejects_stale_latest_without_fabricating_samples(monkeypatch):
    service = AcquisitionService()
    reading = DeviceReading(
        timestamp=datetime.now(UTC),
        raw_power=20,
        raw_power_unit="W",
        power_w=20,
        temperatures_c=[22],
    )
    for device_id in [1, 2]:
        adapter = ClockedStream(device_id, "electrical" if device_id == 1 else "temperature", 1)
        await adapter.connect()
        service.runtimes[device_id] = DeviceRuntime(
            adapter=adapter, latest=reading, source_role=adapter.role
        )

    async def already_connected(device_id):
        return {}

    monkeypatch.setattr(service, "connect", already_connected)
    with pytest.raises(TimeoutError):
        await service.prepare_common_start([1, 2], 1500, timeout_seconds=0.03)
    assert all(
        r.session_id is None and r.pending_start is None and not r.buffer
        for r in service.runtimes.values()
    )
